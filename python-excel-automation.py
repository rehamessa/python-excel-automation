#-------------------------------------------------------
#open the spreedsheets
#-------------------------------------------------------

from openpyxl import load_workbook
workbook=load_workbook('sales.xlsx')
sheet=workbook.active

#-------------------------------------------------------
#Read all sales value
#-------------------------------------------------------
for row in range(2,sheet.max_row+1):
    sales=sheet[f"B{row}" ] .value
    print(sales)

    